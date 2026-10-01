# -*- coding: utf-8 -*-
"""Step 17: 参数化研判全流程（build → judge → xlsx → 溯源关键词候选）。
可独立运行:
    python step17_adjudicate.py --csv <grey_final.csv> --out <输出目录> --work <工作目录>
依赖 work/ 下: step12_build.py(规范化/分组函数), judge_engine.py(知识库+黑名单),
              url_verify.json(网址核验缓存), blacklist_sigs.json(黑名单, 缺省跳过)。
产物: <out>/grey_final_result.xlsx + <out>/kw_candidates.json(溯源待查清单)。
"""
import argparse
import csv
import json
import os
import re
import sys
from collections import Counter, defaultdict

NOISE = {"appdata", "local", "locallow", "roaming", "temp", "tmp", "program files",
         "program files (x86)", "programdata", "common files", "windows", "system32",
         "syswow64", "users", "public", "desktop", "documents", "downloads", "programs",
         "python", "node_modules", "site-packages", "lib", "lib64", "libexec", "binaries",
         "binary", "runtime", "runtimes", "envs", "venv", "packages", "cache", "caches",
         "data", "logs", "log", "resources", "app", "apps", "common", "files", "versions",
         "version", "upgrade", "updates", "update", "autoupdate", "components", "plugins",
         "extensions", "modules", "sbin", "usr", "etc", "opt", "var", "home", "x86", "x64",
         "amd64", "intel", "config", "settings", "prefs", "locales", "i18n", "microsoft",
         "windowsapps", "fonts", "dictionaries", "crashreports", "local settings",
         "application data", "documents and settings", "share", "include", "meta", "sdk",
         "drivers", "store", "games", "game", "server", "client", "db", "database",
         "tools", "software"}
VER_RE = re.compile(r"^(v?\d+([\._]\d+)*|\d{4}|[0-9a-f]{8,}|\d+x\d+)$", re.I)
SIG_HINT = ("ltd", "inc", "llc", "corp", "co.,", "technolog", "company", "software",
            "foundation", "gmbh", "有限", "科技", "股份", "网络", "信息", "软件", "工作室")
HEAD = ["来源", "md5", "路径", "文件名", "签名", "安全指数", "软件名称", "软件网址",
        "软件简短介绍", "备注"]


def dec(s):
    return s.replace("+", " ").strip().lower() if isinstance(s, str) else s


def build_recs(csv_path):
    """复用 step12 的规范化/标志/排序逻辑（保持与既有管线一致）。
    路径中的 '+' 传输编码(代替空格)在规范化前解码, 与判定时 dec() 口径一致,
    否则 program+files 等结构段会逃过 windowsapps/盘符等 flags 识别。"""
    import step12_build as sb
    rows = list(csv.reader(open(csv_path, encoding="utf-8-sig")))
    recs = []
    for r in rows[1:]:
        if len(r) < 5:
            continue
        src, md5, path, fn, sig = r[:5]
        raw = path.strip()
        np_ = sb.norm_path(path.replace("+", " "))
        recs.append({"src": src, "md5": md5, "path": path, "norm": np_, "fn": fn,
                     "sig": sig, "unix": raw.startswith("/") or raw.startswith("~"),
                     "unc": raw.startswith("\\\\"),
                     **sb.flags(np_)})
    recs.sort(key=lambda x: (x["norm"], x["fn"].strip().lower(), x["md5"]))
    return recs


def write_xlsx(judged, out_path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "判定结果"
    ws.append(HEAD)
    hf = Font(name="等线", size=10, bold=True, color="FFFFFF")
    hfill = PatternFill("solid", start_color="4472C4")
    for c in range(1, len(HEAD) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = hf
        cell.fill = hfill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    body = Font(name="等线", size=10)
    for r in judged:
        ws.append([r["src"], r["md5"], r["path"], r["fn"], r["sig"],
                   r["score"], r["name"] or "", r["url"] or "", r["intro"] or "",
                   "AI" if r["ai"] else ""])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = body
            cell.alignment = Alignment(vertical="center")
    for i, w in enumerate([7, 34, 55, 26, 30, 9, 24, 36, 15, 7], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:J%d" % (len(judged) + 1)
    wb.save(out_path)


def _segs(norm):
    """路径中全部有效段（非噪声/非版本/非@环境段，小写）。"""
    out = []
    for s in (norm or "").split("\\"):
        s = s.strip().lower()
        if s and not s.startswith("@") and s not in NOISE \
                and not VER_RE.match(s) and len(s) > 2:
            out.append(s)
    return out


def extract_kw(recs, judged):
    """溯源待查清单: 未定性行按行级关键词聚合, 行数降序=重点溯源优先级。

    厂商目录拆分: 首段下出现 ≥2 个不同产品子段(各≥2行)时, 视为厂商目录,
    关键词取「厂商/子段」复合词(如 tencent/deja), 避免同厂商不同产品混进一个待查词。
    另记录 paths(前两段前缀样例)供过目。"""
    rows_kw = []
    for r, j in zip(recs, judged):
        if j["kind"] != "unknown":
            continue
        rows_kw.append((r, _segs(dec(r["norm"]))))
    sub = defaultdict(Counter)
    for _r, segs in rows_kw:
        if len(segs) >= 2:
            sub[segs[0]][segs[1]] += 1
    vendor = {k for k, c in sub.items()
              if sum(1 for v in c.values() if v >= 2) >= 2}
    kw = defaultdict(lambda: {"rows": 0, "md5": set(), "sig": set(),
                              "fn": set(), "paths": set()})
    for r, segs in rows_kw:
        if segs:
            key = (segs[0] + "/" + segs[1]) if segs[0] in vendor and len(segs) >= 2 \
                else segs[0]
        else:
            fn = dec(r["fn"])
            key = "fn:" + fn if fn and fn != "\\n" else "未识别"
        e = kw[key]
        e["rows"] += 1
        e["md5"].add(r["md5"])
        if segs:
            e["paths"].add("\\".join(segs[:2]))
        sd = dec(r["sig"])
        if sd and sd != "\\n" and any(h in sd for h in SIG_HINT):
            e["sig"].add(sd)
        fn = dec(r["fn"])
        if fn and fn != "\\n":
            e["fn"].add(fn)
    out = [{"kw": k, "rows": v["rows"], "md5": len(v["md5"]),
            "sig": sorted(v["sig"])[:3], "fn": sorted(v["fn"])[:5],
            "paths": sorted(v["paths"])[:5]}
           for k, v in kw.items()]
    out.sort(key=lambda x: -x["rows"])
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--csv", required=True, help="grey_final.csv 路径")
    ap.add_argument("--out", required=True, help="结果输出目录")
    ap.add_argument("--work", required=True, help="工作目录(规则引擎/黑名单/网址核验)")
    a = ap.parse_args()
    sys.path.insert(0, a.work)
    import judge_engine as je

    recs = build_recs(a.csv)
    json.dump(recs, open(os.path.join(a.work, "sorted_rows.json"), "w", encoding="utf-8"),
              ensure_ascii=False)

    vf = os.path.join(a.work, "url_verify.json")
    url_ok = json.load(open(vf, encoding="utf-8")) if os.path.exists(vf) else {}
    bf = os.path.join(a.work, "blacklist_sigs.json")
    blacklist = [e["norm"] for e in json.load(open(bf, encoding="utf-8"))] \
        if os.path.exists(bf) else []
    judged = je.judge_all(recs, url_ok, blacklist)
    json.dump(judged, open(os.path.join(a.work, "judged_rows.json"), "w", encoding="utf-8"),
              ensure_ascii=False)

    xlsx = os.path.join(a.out, "grey_final_result.xlsx")
    write_xlsx(judged, xlsx)

    kws = extract_kw(recs, judged)
    json.dump(kws, open(os.path.join(a.out, "kw_candidates.json"), "w", encoding="utf-8"),
              ensure_ascii=False, indent=1)

    print("ADJUDICATED rows=%d -> %s" % (len(judged), xlsx))
    print("KINDS", json.dumps(dict(Counter(r["kind"] for r in judged)), ensure_ascii=False))
    print("SCORE_DIST", json.dumps(dict(Counter(str(r["score"]) for r in judged)),
                                   ensure_ascii=False))
    print("HEAD_URL", sum(1 for r in judged if r["head"] and r["url"]),
          "AI", sum(1 for r in judged if r["ai"]))
    print("KW_BACKLOG kw=%d unsettled_rows=%d top5=%s"
          % (len(kws), sum(x["rows"] for x in kws),
             json.dumps([(x["kw"], x["rows"]) for x in kws[:5]], ensure_ascii=False)))


if __name__ == "__main__":
    main()
